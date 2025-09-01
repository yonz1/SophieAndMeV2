import openpyxl as openpyxl
import sqlite3
import re


# Charger le fichier Excel avec openpyxl

con_i = sqlite3.connect("C:\\Users\\Bastien\\RiderProjects\\SophieAndMe\\SophieAndMe\\Database\\EDT.db")
cur_i = con_i.cursor()


ListMat = ["Maths","SI","TIPE","Francais","Sport","Info","Physique","Anglais"]
ListRot = ["Physique","TP SI","TIPE"]
wb_obj = openpyxl.load_workbook('Livre 3 1.xlsx')
sheet_obj = wb_obj.active
Row = sheet_obj.max_row
Column = sheet_obj.max_column
Header = []
val = []
Temp = ""
G = "A"


def FillSalle(data):
    ret = []
    L = re.split(r'[/\s$]+',data)
    for i in range(len(L)):
        if re.fullmatch(r"[A-Za-z]\d{3}",L[i]) or re.fullmatch(r"\d{2}[A-Za-z]",L[i]) :
            ret.append(L[i])
        elif i+1 < len(L) and re.fullmatch(r"[A-Za-z]\d{3}",L[i]+L[i+1]):
            ret.append(L[i]+L[i+1])

    return ret



def GetPos(L):
    for i in range (len(L)):
        if L[i] == "TP" or L[i] == "TD":
            return i


def SearchMat(data):
    if ("TP" in data or "TD" in data):
        L = data.split(" ")
        return L[GetPos(L)] + " " + L[GetPos(L) + 1]
    else:
        for i in ListMat:
            if i in data or i[:3] in data:
                return i

def Multiple(data):
    if (" GA" in data and " GB" in data):
        return True
    return False

def FillGroupe(data):
    ret = []
    val1 = " GA"
    data = data.replace("GAILLY","")
    val2 = " GB"
    for i in range(len(data)):
        if (data[i:i+3] == val1 or data[i:i+3] == "$GA"):
            ret.append("GA")
        elif (data[i:i+3] == val2 or data[i:i+3] == "$GB"):
            ret.append("GB")
    return ret

def FillEnsei(data):
    if ("$" in data):
        L = data.split("$")
        return L[1]
    return ""

def GetInfos(data):
    _Mat = []
    _Enseignant = []
    _Salle = []
    _Groupe = []
    data.replace("CE","")
    _Groupe = FillGroupe(data)
    _Enseignant = FillEnsei(data)
    _Salle = FillSalle(data)
    if Multiple(data):
        val = SearchMat(data)
        _Mat.append(val)
        data = data.replace(val,"")
        _Mat.append(SearchMat(data))
    else:
        _Mat.append(SearchMat(data))
    return (_Mat,_Salle,_Enseignant,_Groupe)




sheet_obj["A3"] = "Heure"
print(sheet_obj.merged_cells)
for merged_range in (list(sheet_obj.merged_cells.ranges)):
    sheet_obj.unmerge_cells((str(merged_range)))
    L = str(merged_range).split(":")
    min_col, min_row, max_col, max_row = merged_range.bounds
    value = sheet_obj.cell(row=min_row, column=min_col).value
    if max_row-min_row > 1:
        for n in range(min_row+1,max_row):
            L.append(L[0][0]+str(n))
    for i in L:
        print(i)
        sheet_obj[i] = value

def GetVal(Groupe,G):
    for p in range(len(Groupe)):
        if Groupe[p] == G:
            return p
    return -1


############################################################### Affichage EDT Groupe A


print("SemaineA")
for i in range(2,Column + 1):
    Mat = []
    Salle = []
    Enseignant = []
    Groupe = []
    Temp = sheet_obj.cell(row=3,column=i).value
    Header.append(val)
    Table = Temp+ "A"
    command = ("DELETE FROM " + Table)
    cur_i.execute(command)
    print(Temp)
    for y in range(4, Row + 1):
        data = str(sheet_obj.cell(row=y,column=i).value).replace("\n","$")
        (Mat,Salle,Enseignant,Groupe) = GetInfos(data)
        if "GA" in Groupe:
            z = GetVal(Groupe,"GA")
            Presence = True
        else:
            Presence = False
        if str(Mat[0]) == "None":
            # print(sheet_obj.cell(row=y, column=1).value +  " - Rien")
            Command = ("INSERT INTO " + Table + " (Mat,Salle,Enseignant) VALUES (?,?,?)")
            val = ("","","")
            print(Command)
            print(val)
            cur_i.execute(Command,val)
            con_i.commit()
        elif (len(Groupe) == 0):
            mat_val = Mat[0]  if Mat else ""
            salle_val = Salle[0]   if Salle else ""
            Enseignant_val = Enseignant if Enseignant else ""
            # print(sheet_obj.cell(row=y, column=1).value + mat_val + salle_val + Enseignant_val)
            Command = ("INSERT INTO " + Table + " (Mat,Salle,Enseignant) VALUES (?,?,?)")
            val = (mat_val,salle_val,Enseignant_val)
            print(Command)
            print(val)
            cur_i.execute(Command,val)
            con_i.commit()
        elif Presence:
            mat_val = Mat[z] if Mat[z] else ""
            salle_val = Salle[z] if Salle[z] else ""
            # print(sheet_obj.cell(row=y, column=1).value + " - " + str(Mat[z]) + " - " + str(Salle[z]) + " - " +Enseignant)
            Command = ("INSERT INTO " + Table + " (Mat,Salle,Enseignant) VALUES (?,?,?)")
            val = (str(Mat[z]),str(Salle[z]),Enseignant)
            print(Command)
            print(val)
            cur_i.execute(Command,val)
            con_i.commit()
        else:
            # print(sheet_obj.cell(row=y, column=1).value +  " - Rien")
            Command = ("INSERT INTO " + Table + " (Mat,Salle,Enseignant) VALUES (?,?,?)")
            val = ("","","")
            print(Command)
            print(val)
            cur_i.execute(Command,val)
            con_i.commit()
        print("############################################")


############################################################### Affichage EDT Groupe B
print("############################################")
print("############################################")
print("############################################")

print("SemaineB")
for i in range(2,Column + 1):
    Mat = []
    Salle = []
    Enseignant = []
    Groupe = []
    Temp = sheet_obj.cell(row=3,column=i).value
    Table = Temp+ "B"
    command = ("DELETE FROM " + Table)
    cur_i.execute(command)
    Header.append(val)
    print(Temp)
    for y in range(4, Row + 1):
        data = str(sheet_obj.cell(row=y,column=i).value).replace("\n","$")
        (Mat,Salle,Enseignant,Groupe) = GetInfos(data)
        if "GB" in Groupe:
            z = GetVal(Groupe,"GB")
            Presence = True
        else:
            Presence = False
        if str(Mat[0]) == "None":
            # print(sheet_obj.cell(row=y, column=1).value +  " - Rien")
            Command = ("INSERT INTO " + Table + " (Mat,Salle,Enseignant) VALUES (?,?,?)")
            val = ("","","")
            print(Command)
            print(val)
            cur_i.execute(Command,val)
            con_i.commit()
        elif (len(Groupe) == 0):
            mat_val = Mat[0]  if Mat else ""
            salle_val = Salle[0]   if Salle else ""
            Enseignant_val = Enseignant if Enseignant else ""
            # print(sheet_obj.cell(row=y, column=1).value + mat_val + salle_val + Enseignant_val)
            Command = ("INSERT INTO " + Table + " (Mat,Salle,Enseignant) VALUES (?,?,?)")
            val = (mat_val,salle_val,Enseignant_val)
            print(Command)
            print(val)
            cur_i.execute(Command,val)
            con_i.commit()
        elif Presence:
            if(Mat in ListRot):
                mat_val = Mat[z] if Mat[z] else ""
                salle_val = Salle[z] if Salle[z] else ""
                # print(sheet_obj.cell(row=y, column=1).value + " - " + str(Mat[z]) + " - " + str(Salle[z]) + " - " +Enseignant)
                Command = ("INSERT INTO " + Table + " (Mat,Salle,Enseignant) VALUES (?,?,?)")
                val = (str(Mat[z]),str(Salle[z]),Enseignant)
                print(Command)
                print(val)
                cur_i.execute(Command,val)
                con_i.commit()
            else:
                z = GetVal(Groupe, "GA")
                mat_val = Mat[z] if Mat[z] else ""
                salle_val = Salle[z] if Salle[z] else ""
                # print(sheet_obj.cell(row=y, column=1).value + " - " + str(Mat[z]) + " - " + str(Salle[z]) + " - " +Enseignant)
                Command = ("INSERT INTO " + Table + " (Mat,Salle,Enseignant) VALUES (?,?,?)")
                val = (str(Mat[z]),str(Salle[z]),Enseignant)
                print(Command)
                print(val)
                cur_i.execute(Command,val)
                con_i.commit()
        else:
            # print(sheet_obj.cell(row=y, column=1).value +  " - Rien")
            Command = ("INSERT INTO " + Table + " (Mat,Salle,Enseignant) VALUES (?,?,?)")
            val = ("","","")
            print(Command)
            print(val)
            cur_i.execute(Command,val)
            con_i.commit()
        print("############################################")


